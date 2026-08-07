#!/usr/bin/env python3
"""Export validated, globally deduplicated Xiaohongshu JSON files to Excel."""

import datetime
import glob
import json
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from xhs_export_utils import deduplicate_posts

ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
HEADERS = [
    "话题",
    "序号",
    "note_id",
    "标题",
    "作者",
    "日期",
    "正文",
    "标签",
    "互动",
    "评论数",
    "评论(前8)",
    "图片数",
    "链接",
]


def clean(value):
    return ILLEGAL.sub("", value) if isinstance(value, str) else value


def load_records(base, date, folder_tag, slugs):
    records = []
    issues = []
    for slug in slugs:
        directory = os.path.join(base, f"{date}_{folder_tag}_{slug}")
        files = [
            path
            for path in glob.glob(os.path.join(directory, "*.json"))
            if os.path.basename(path) != "state.json"
        ]
        for path in sorted(files):
            try:
                with open(path, encoding="utf-8") as handle:
                    note = json.load(handle)
                if not isinstance(note, dict):
                    raise TypeError("root JSON value is not an object")
                records.append({"slug": slug, "path": path, "note": note})
            except (OSError, TypeError, ValueError, json.JSONDecodeError) as error:
                issues.append(
                    {
                        "slug": slug,
                        "path": path,
                        "note": {},
                        "reason": "invalid-json",
                        "detail": str(error),
                        "actual_id": "",
                    }
                )
    return records, issues


def initialize_workbook():
    workbook = Workbook()
    summary = workbook.active
    summary.title = "汇总"
    summary.append(
        [
            "话题",
            "有效篇数",
            "评论数",
            "图片数",
            "均互动(数1/数2/数3)",
            "跳过重复",
            "跳过无效",
        ]
    )
    for cell in summary[1]:
        cell.font = Font(bold=True)

    posts = workbook.create_sheet("全部帖子")
    posts.append(HEADERS)
    for cell in posts[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(
            start_color="D6EAF8", end_color="D6EAF8", fill_type="solid"
        )
    widths = [10, 5, 26, 32, 14, 12, 60, 22, 16, 7, 50, 7, 48]
    for index, width in enumerate(widths, 1):
        posts.column_dimensions[get_column_letter(index)].width = width
    posts.freeze_panes = "A2"

    quality = workbook.create_sheet("质量报告")
    quality.append(["话题", "文件", "保存ID", "真实ID", "原因", "详情"])
    for cell in quality[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(
            start_color="FADBD8", end_color="FADBD8", fill_type="solid"
        )
    quality.column_dimensions["A"].width = 14
    quality.column_dimensions["B"].width = 60
    quality.column_dimensions["C"].width = 26
    quality.column_dimensions["D"].width = 26
    quality.column_dimensions["E"].width = 28
    quality.column_dimensions["F"].width = 60
    quality.freeze_panes = "A2"
    return workbook, summary, posts, quality


def export(base, date, folder_tag, slugs, output_tag=None):
    records, load_issues = load_records(base, date, folder_tag, slugs)
    unique, data_issues = deduplicate_posts(records)
    issues = load_issues + data_issues
    workbook, summary, posts, quality = initialize_workbook()
    thin = Border(bottom=Side(style="thin", color="DDDDDD"))

    stats = {
        slug: {
            "posts": 0,
            "comments": 0,
            "images": 0,
            "engagement": [0, 0, 0],
            "engagement_rows": 0,
            "duplicates": 0,
            "invalid": 0,
        }
        for slug in slugs
    }

    per_slug_index = {slug: 0 for slug in slugs}
    for record in unique:
        slug = record["slug"]
        note = record["note"]
        per_slug_index[slug] += 1
        stat = stats[slug]
        comments = note.get("comments") or []
        images = note.get("images") or []
        engagement = note.get("engageNums") or []
        stat["posts"] += 1
        stat["comments"] += len(comments)
        stat["images"] += len(images)
        if len(engagement) >= 3:
            for index in range(3):
                stat["engagement"][index] += engagement[index]
            stat["engagement_rows"] += 1

        top_comments = " | ".join(
            f"{(comment.get('name') or '').strip()}:{clean(comment.get('content') or '').strip()[:60]}"
            for comment in comments[:8]
        )
        posts.append(
            [
                slug,
                per_slug_index[slug],
                note.get("id", ""),
                clean(note.get("title", "")),
                clean(note.get("author", "")),
                clean(note.get("date", "")),
                clean(note.get("desc", "")),
                clean(" ".join(note.get("tags") or [])),
                "/".join(str(value) for value in engagement[:4]),
                len(comments),
                clean(top_comments),
                len(images),
                note.get("url", ""),
            ]
        )
        for cell in posts[posts.max_row]:
            cell.border = thin
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=cell.column in (4, 7, 8, 11),
            )

    for issue in issues:
        reason = issue["reason"]
        slug = issue.get("slug", "")
        if slug in stats:
            key = "duplicates" if reason.startswith("duplicate-") else "invalid"
            stats[slug][key] += 1
        note = issue.get("note") or {}
        quality.append(
            [
                slug,
                issue.get("path", ""),
                note.get("id", ""),
                issue.get("actual_id", ""),
                reason,
                issue.get("detail", ""),
            ]
        )

    for slug in slugs:
        stat = stats[slug]
        count = stat["engagement_rows"]
        average = (
            [value // count for value in stat["engagement"]] if count else [0, 0, 0]
        )
        summary.append(
            [
                slug,
                stat["posts"],
                stat["comments"],
                stat["images"],
                "/".join(map(str, average)),
                stat["duplicates"],
                stat["invalid"],
            ]
        )

    totals = {
        key: sum(stat[key] for stat in stats.values())
        for key in ("posts", "comments", "images", "duplicates", "invalid")
    }
    summary.append(
        [
            "合计",
            totals["posts"],
            totals["comments"],
            totals["images"],
            "",
            totals["duplicates"],
            totals["invalid"],
        ]
    )
    for cell in summary[summary.max_row]:
        cell.font = Font(bold=True)

    if output_tag is None:
        output_tag = "_全" if len(slugs) > 6 else "_" + "_".join(slugs)
    os.makedirs(base, exist_ok=True)
    output = os.path.join(base, f"{date}_{folder_tag}_汇总{output_tag}.xlsx")
    workbook.save(output)
    return output, totals


def main(argv=None):
    args = list(sys.argv[1:] if argv is None else argv)
    if not args:
        print("用法: python3 xhs_combined_export.py <slug1> <slug2> ...")
        return 2

    base = os.environ.get("XHS_DATA_DIR", os.path.join(os.getcwd(), "xhs_data"))
    date = os.environ.get(
        "XHS_DATE", datetime.datetime.now().astimezone().strftime("%Y%m%d")
    )
    folder_tag = os.environ.get("XHS_FOLDERTAG", "xhs")
    output, totals = export(
        base,
        date,
        folder_tag,
        args,
        os.environ.get("XHS_OUTTAG"),
    )
    print(
        f"✅ 汇总: {totals['posts']} 篇 / {totals['comments']} 评论 / "
        f"{totals['images']} 图 → {output}"
    )
    if totals["duplicates"] or totals["invalid"]:
        print(
            f"⚠️ 质量过滤: 跳过 {totals['duplicates']} 个重复、"
            f"{totals['invalid']} 个无效文件；详情见 Excel 的「质量报告」sheet"
        )
    print(f"   大小: {os.path.getsize(output) // 1024} KB")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
